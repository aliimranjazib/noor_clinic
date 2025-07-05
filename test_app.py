#!/usr/bin/env python3
"""
Test script for Radiologist Report Generator
"""

import os
import sys
import tempfile
import shutil
from datetime import datetime

def test_imports():
    """Test if all required modules can be imported"""
    print("Testing imports...")
    
    try:
        import tkinter as tk
        print("✓ tkinter imported successfully")
    except ImportError as e:
        print(f"✗ Failed to import tkinter: {e}")
        return False
    
    try:
        from docx import Document
        print("✓ python-docx imported successfully")
    except ImportError as e:
        print(f"✗ Failed to import python-docx: {e}")
        return False
    
    try:
        from docx.shared import Inches
        from docx.enum.text import WD_ALIGN_PARAGRAPH
        print("✓ docx submodules imported successfully")
    except ImportError as e:
        print(f"✗ Failed to import docx submodules: {e}")
        return False
    
    return True

def test_document_creation():
    """Test if Word document creation works"""
    print("\nTesting Word document creation...")
    
    try:
        from docx import Document
        from docx.enum.text import WD_ALIGN_PARAGRAPH
        
        # Create a test document
        doc = Document()
        
        # Add content
        title = doc.add_heading('Test Radiology Report', 0)
        title.alignment = WD_ALIGN_PARAGRAPH.CENTER
        
        doc.add_paragraph("Radiologist: Dr. Test Johnson")
        doc.add_paragraph(f"Date: {datetime.now().strftime('%B %d, %Y')}")
        
        doc.add_heading('Patient Information', level=1)
        doc.add_paragraph("Name: Test Patient")
        doc.add_paragraph("Age: 45")
        doc.add_paragraph("Test Type: X-Ray")
        
        doc.add_heading('Diagnosis / Report', level=1)
        doc.add_paragraph("This is a test report content.")
        
        # Save to temporary file
        with tempfile.NamedTemporaryFile(suffix='.docx', delete=False) as tmp:
            doc.save(tmp.name)
            tmp_path = tmp.name
        
        # Check if file was created
        if os.path.exists(tmp_path):
            file_size = os.path.getsize(tmp_path)
            print(f"✓ Test document created successfully ({file_size} bytes)")
            
            # Clean up
            os.unlink(tmp_path)
            return True
        else:
            print("✗ Test document was not created")
            return False
            
    except Exception as e:
        print(f"✗ Failed to create test document: {e}")
        return False

def test_reports_folder():
    """Test if Reports folder can be created"""
    print("\nTesting Reports folder creation...")
    
    test_folder = "test_reports"
    
    try:
        if not os.path.exists(test_folder):
            os.makedirs(test_folder)
            print(f"✓ Created test folder: {test_folder}")
        
        # Test file creation in folder
        test_file = os.path.join(test_folder, "test.txt")
        with open(test_file, 'w') as f:
            f.write("Test content")
        
        if os.path.exists(test_file):
            print("✓ Test file created in folder")
            
            # Clean up
            os.unlink(test_file)
            shutil.rmtree(test_folder)
            print("✓ Cleaned up test folder")
            return True
        else:
            print("✗ Failed to create test file")
            return False
            
    except Exception as e:
        print(f"✗ Failed to test folder creation: {e}")
        return False

def test_app_class():
    """Test if the RadiologistApp class can be instantiated"""
    print("\nTesting RadiologistApp class...")
    
    try:
        # Import the app class
        sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
        from radiologist_app import RadiologistApp
        
        print("✓ RadiologistApp class imported successfully")
        
        # Test if we can create an instance (without showing GUI)
        import tkinter as tk
        root = tk.Tk()
        root.withdraw()  # Hide the window
        
        app = RadiologistApp(root)
        print("✓ RadiologistApp instance created successfully")
        
        # Test basic attributes
        if hasattr(app, 'radiologist_name'):
            print(f"✓ Radiologist name: {app.radiologist_name}")
        
        if hasattr(app, 'patient_name_entry'):
            print("✓ Patient name entry field exists")
        
        if hasattr(app, 'save_report'):
            print("✓ Save report method exists")
        
        # Clean up
        root.destroy()
        return True
        
    except Exception as e:
        print(f"✗ Failed to test RadiologistApp class: {e}")
        return False

def main():
    """Run all tests"""
    print("=" * 50)
    print("Radiologist Report Generator - Test Suite")
    print("=" * 50)
    
    tests = [
        ("Import Test", test_imports),
        ("Document Creation Test", test_document_creation),
        ("Reports Folder Test", test_reports_folder),
        ("App Class Test", test_app_class)
    ]
    
    passed = 0
    total = len(tests)
    
    for test_name, test_func in tests:
        print(f"\n--- {test_name} ---")
        if test_func():
            passed += 1
        else:
            print(f"✗ {test_name} FAILED")
    
    print("\n" + "=" * 50)
    print(f"Test Results: {passed}/{total} tests passed")
    
    if passed == total:
        print("🎉 All tests passed! The application should work correctly.")
        return True
    else:
        print("⚠️  Some tests failed. Please check the errors above.")
        return False

if __name__ == "__main__":
    success = main()
    sys.exit(0 if success else 1) 